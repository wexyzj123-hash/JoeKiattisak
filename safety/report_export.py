"""safety/report_export.py — ส่งออกรายงาน (Report) เป็น Excel (openpyxl) และ PDF (reportlab + ฟอนต์ไทยที่ฝังในที่เก็บ)"""
from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table as RlTable, TableStyle

from . import report_fonts
from .reports import Report

CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
INVALID_SHEET = re.compile(r"[\[\]:*?/\\]")
BRAND = "1F4E8C"


def safe_cell(value):
    """ข้อความที่ขึ้นต้นด้วย = + - @ ถูกเขียนเป็นข้อความ (กัน formula injection ใน Excel)"""
    if value is None:
        return ""
    if isinstance(value, str):
        value = CONTROL_RE.sub("", value)                       # openpyxl ปฏิเสธอักขระควบคุม — ตัดทิ้งแทนที่จะทำให้ส่งออกล้ม
        if value[:1] in ("=", "+", "-", "@", "\t", "\r"):
            return "'" + value
    return value


def _sheet_name(title: str, used: set[str]) -> str:
    base = (INVALID_SHEET.sub(" ", title).strip() or "ตาราง")[:31]
    name, n = base, 2
    while name.lower() in used:
        suffix = f" {n}"
        name, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(name.lower())
    return name


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "สรุป"
    used = {"สรุป"}
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=BRAND)
    ws["A1"], ws["A1"].font = safe_cell(report.title), Font(bold=True, size=14)
    ws["A2"] = safe_cell(report.subtitle)
    ws["A3"] = f"งวด: {report.period_label}"
    ws["A4"] = f"จัดทำเมื่อ: {report.generated_at}"
    row = 6
    for k in report.kpis:
        ws.cell(row=row, column=1, value=safe_cell(k.label))
        ws.cell(row=row, column=2, value=safe_cell(k.value))
        ws.cell(row=row, column=3, value=safe_cell(k.note))
        row += 1
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 42, 18, 36
    for t in report.tables:
        sheet = wb.create_sheet(_sheet_name(t.title, used))
        for c, col in enumerate(t.columns, 1):
            cell = sheet.cell(row=1, column=c, value=safe_cell(col))
            cell.font, cell.fill, cell.alignment = head, fill, Alignment(wrap_text=True, vertical="center")
        for r, data in enumerate(t.rows, 2):
            for c, v in enumerate(data, 1):
                sheet.cell(row=r, column=c, value=safe_cell(v)).alignment = Alignment(wrap_text=True, vertical="top")
        if not t.rows:
            sheet.cell(row=2, column=1, value=t.empty_text)
        for c, col in enumerate(t.columns, 1):
            longest = max([len(str(col))] + [len(str(r[c - 1])) for r in t.rows if c - 1 < len(r)])
            sheet.column_dimensions[get_column_letter(c)].width = min(60, max(12, longest + 2))
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _Numbered(rl_canvas.Canvas):
    def __init__(self, *a, font="ReportFont", **kw):
        super().__init__(*a, **kw)
        self._saved, self._font = [], font

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved)
        for state in self._saved:
            self.__dict__.update(state)
            self.setFont(self._font, 8)
            self.setFillColor(colors.grey)
            self.drawRightString(A4[0] - 15 * mm, 10 * mm, f"หน้า {self._pageNumber}/{total}")
            super().showPage()
        super().save()


def to_pdf(report: Report) -> bytes:
    regular, bold = report_fonts.register()
    p = lambda text, **kw: Paragraph(escape(str(text)), ParagraphStyle("p", fontName=kw.pop("font", regular), leading=kw.get("fontSize", 9) * 1.35, **{"fontSize": 9, **kw}))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=18 * mm,
                            title=report.title, author="Safety 4.0")
    story = [p(report.title, font=bold, fontSize=17, textColor=colors.HexColor("#" + BRAND)), p(report.subtitle, fontSize=10),
             p(f"งวด: {report.period_label}   ·   จัดทำเมื่อ {report.generated_at}", fontSize=8, textColor=colors.grey), Spacer(1, 6 * mm)]
    width = A4[0] - 30 * mm
    if report.kpis:
        cells = [[p(k.label, fontSize=8, textColor=colors.HexColor("#4a5a70")), p(k.value, font=bold, fontSize=14 if isinstance(k.value, (int, float)) else 9)] for k in report.kpis]
        per = 3
        grid = [sum((cells[i + j] if i + j < len(cells) else ["", ""] for j in range(per)), []) for i in range(0, len(cells), per)]
        t = RlTable(grid, colWidths=[width / (per * 2) * (1.4 if c % 2 == 0 else 0.6) for c in range(per * 2)])
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#c8d3e2")), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e2e8f2")),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        story += [t, Spacer(1, 6 * mm)]
    for tab in report.tables:
        story.append(p(tab.title, font=bold, fontSize=11))
        story.append(Spacer(1, 2 * mm))
        head = [p(c, font=bold, fontSize=8, textColor=colors.white) for c in tab.columns]
        body = [[p(v, fontSize=8) for v in row] for row in tab.rows] or [[p(tab.empty_text, fontSize=8, textColor=colors.grey)] + [""] * (len(tab.columns) - 1)]
        t = RlTable([head] + body, colWidths=[width / len(tab.columns)] * len(tab.columns), repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BRAND)), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c8d3e2")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f7fb")])]))
        story += [t, Spacer(1, 6 * mm)]
    doc.build(story, canvasmaker=lambda *a, **kw: _Numbered(*a, font=regular, **kw))
    return buf.getvalue()
