import io
import re

import pytest
from openpyxl import load_workbook

from safety import report_export as ex
from safety.reports import Kpi, Report, Table


def rep(**kw):
    base = dict(key="monthly", title="รายงานประจำเดือน", subtitle="โรงงาน 1 / แผนก ก", period_label="มีนาคม 2569", generated_at="20/03/2026 10:00",
                kpis=[Kpi("แจ้งอันตราย", 5), Kpi("วันปลอดอุบัติเหตุ", "ไม่มีข้อมูล", "หมายเหตุ")],
                tables=[Table("ตารางที่ 1", ["รหัส", "หัวข้อ"], [["HZ-1", "เครื่องจักรไม่มีการ์ด"], ["HZ-2", "ก้ำก่ำ น้ำท่วม ผู้ชื่นชม"]]),
                        Table("ตารางว่าง", ["ก", "ข"], [])])
    base.update(kw)
    return Report(**base)


def pages(pdf: bytes) -> int:
    return len(re.findall(rb"/Type\s*/Page[^s]", pdf))


def test_xlsx_roundtrip_values_and_sheets():
    wb = load_workbook(io.BytesIO(ex.to_xlsx(rep())))
    assert wb.sheetnames[0] == "สรุป" and len(wb.sheetnames) == 3
    first = wb["สรุป"]
    assert first["A1"].value == "รายงานประจำเดือน" and ("แจ้งอันตราย", 5) in [(r[0].value, r[1].value) for r in first.iter_rows()]
    t1 = wb[wb.sheetnames[1]]
    assert [c.value for c in t1[1]] == ["รหัส", "หัวข้อ"] and t1["B3"].value == "ก้ำก่ำ น้ำท่วม ผู้ชื่นชม" and t1.freeze_panes == "A2"


def test_xlsx_formula_injection_is_text():
    bad = '=HYPERLINK("http://x","y")'
    wb = load_workbook(io.BytesIO(ex.to_xlsx(rep(tables=[Table("t", ["a"], [[bad], ["+1"], ["@x"], ["-2"]])]))))
    ws = wb[wb.sheetnames[1]]
    for row in range(2, 6):
        cell = ws.cell(row=row, column=1)
        assert cell.data_type == "s" and cell.value.startswith("'")
    assert ex.safe_cell(5) == 5 and ex.safe_cell(None) == ""


def test_xlsx_sheet_names_valid_and_unique():
    tabs = [Table("ชื่อ/ยาว[มาก]:*?" + "ก" * 60, ["a"], [[1]]) for _ in range(3)]
    names = load_workbook(io.BytesIO(ex.to_xlsx(rep(tables=tabs)))).sheetnames
    assert len(set(names)) == len(names) == 4 and all(len(n) <= 31 and not re.search(r"[\[\]:*?/\\]", n) for n in names)


def test_pdf_is_valid_and_embeds_thai_font():
    pdf = ex.to_pdf(rep())
    assert pdf.startswith(b"%PDF") and b"/FontFile2" in pdf and pages(pdf) == 1


def test_pdf_escapes_markup_in_cells():
    pdf = ex.to_pdf(rep(tables=[Table("<b>t</b>", ["<i>a</i>"], [["<b>&</b> </para>"]])], title="<x>"))
    assert pdf.startswith(b"%PDF")


def test_pdf_paginates_300_rows_and_long_text():
    rows = [[f"HZ-{i}", "ข้อความยาวมาก " * 20] for i in range(300)]
    assert pages(ex.to_pdf(rep(tables=[Table("ใหญ่", ["รหัส", "หัวข้อ"], rows)]))) > 3


def test_empty_report_exports_both():
    empty = rep(kpis=[], tables=[])
    assert load_workbook(io.BytesIO(ex.to_xlsx(empty))).sheetnames == ["สรุป"] and ex.to_pdf(empty).startswith(b"%PDF")


def test_xlsx_control_characters_and_tab_cr_are_neutralised():
    import openpyxl
    wb = load_workbook(io.BytesIO(ex.to_xlsx(rep(tables=[Table("t", ["a"], [["\x01bad\x0b"], ["\tx"], ["\ry"]])]))))
    ws = wb[wb.sheetnames[1]]
    assert ws["A2"].value == "bad" and ws["A3"].value.startswith("'") and ws["A4"].value.startswith("'")
