import io

import pytest
from openpyxl import load_workbook

from safety import timeutil
from safety.db import connect
from tests.safety.factories import add_chain, add_hazard
from tests.safety.helpers import login, make_user


@pytest.fixture
def env(safety_app):
    con = connect(safety_app.config["DATABASE"])
    a, b = add_chain(con, "A"), add_chain(con, "B")
    admin = login(safety_app, "admin", name="แอดมิน")
    add_hazard(con, a["a"], admin.uid, reported_at="2026-03-02 08:00")
    add_hazard(con, b["a"], admin.uid, reported_at="2026-03-02 08:00")
    con.close()
    head = login(safety_app, "dept_head", name="หัวหน้า", factory_id=a["f"])
    rep = login(safety_app, "reporter", name="ผู้แจ้ง")
    viewer = login(safety_app, "viewer", name="ผู้ดู")
    con = connect(safety_app.config["DATABASE"])
    con.execute("UPDATE safety_roles SET permissions='master.view,report.view' WHERE key='viewer'")
    con.commit(); con.close()
    return type("E", (), dict(app=safety_app, a=a, b=b, admin=admin, head=head, rep=rep, viewer=viewer))


def test_index_requires_report_view(env):
    assert env.rep.get("/reports").status_code == 403
    r = env.head.get("/reports")
    assert r.status_code == 200 and "รายงานประจำเดือนตามแผนก" in r.get_data(as_text=True)


def test_view_renders_kpis_and_tables(env):
    page = env.admin.get("/reports/monthly?year=2026&month=3").get_data(as_text=True)
    assert "มีนาคม 2569" in page and "แจ้งอันตราย" in page and "อันตรายที่แจ้งแยกตามประเภท" in page
    assert env.admin.get("/reports/nope").status_code == 404


def test_scoped_view_excludes_other_factory(env):
    page = env.head.get("/reports/monthly?year=2026&month=3").get_data(as_text=True)
    assert "พื้นที่ A" in page and "พื้นที่ B" not in page


def test_export_needs_report_export_and_returns_valid_files(env):
    assert env.viewer.get("/reports/monthly/xlsx?year=2026&month=3").status_code == 403
    r = env.head.get("/reports/monthly/xlsx?year=2026&month=3")
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    assert load_workbook(io.BytesIO(r.data)).sheetnames[0] == "สรุป"
    p = env.admin.get("/reports/annual/pdf?year=2026")
    assert p.status_code == 200 and p.data.startswith(b"%PDF") and p.mimetype == "application/pdf"
    assert env.admin.get("/reports/capa/xlsx").status_code == 200 and env.admin.get("/reports/monthly/zip").status_code == 404


def test_out_of_scope_department_404(env):
    assert env.head.get(f"/reports/monthly?department_id={env.b['d']}").status_code == 404
    assert env.head.get(f"/reports/monthly/pdf?department_id={env.b['d']}").status_code == 404
    assert env.head.get(f"/reports/monthly?department_id={env.a['d']}&year=2026&month=3").status_code == 200


def test_filename_header_is_ascii_safe_and_bad_period_falls_back(env):
    r = env.admin.get("/reports/monthly/xlsx?year=2026&month=3")
    assert r.headers["Content-Disposition"] == 'attachment; filename="report-monthly-2026-03.xlsx"'
    page = env.admin.get("/reports/monthly?year=1&month=99", follow_redirects=True).get_data(as_text=True)
    assert "ปีหรือเดือนไม่ถูกต้อง" in page


def test_pdf_without_font_redirects_with_message(env, monkeypatch):
    from safety import report_fonts
    monkeypatch.setattr(report_fonts, "FONT_DIR", "/nonexistent")
    r = env.admin.get("/reports/annual/pdf?year=2026")
    assert r.status_code == 302 and "/reports/annual" in r.headers["Location"]
    assert env.admin.get("/reports/annual/xlsx?year=2026").status_code == 200


def test_new_permissions_apply_to_default_roles_only_for_new_installs(safety_app):
    from safety.perms import role_permissions
    con = connect(safety_app.config["DATABASE"])
    assert {"report.view", "report.export", "report.manage"} <= role_permissions(con, "assessor")
    assert {"report.view", "report.export"} <= role_permissions(con, "dept_head") and "report.manage" not in role_permissions(con, "dept_head")
    assert not {"report.view", "report.export"} & role_permissions(con, "viewer")             # แอดมินให้สิทธิ์เองที่ /admin/roles
    assert "report.view" not in role_permissions(con, "reporter")
